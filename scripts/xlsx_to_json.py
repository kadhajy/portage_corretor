"""Converte tabela_portage.xlsx nos dados usados pela aplicação.

Gera `data/portage.json`, que a página carrega via fetch.

Utilitário opcional (única dependência Python do projeto). A planilha
`tabela_portage.xlsx` não faz parte do repositório. Para regenerar,
coloque-a na raiz (cada planilha vira uma "área"; linhas sem conteúdo na
coluna `habilidade` são ignoradas) e rode a partir da raiz:

    pip install openpyxl
    python scripts/xlsx_to_json.py
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
XLSX_FILE = BASE_DIR / "tabela_portage.xlsx"
JSON_FILE = BASE_DIR / "data" / "portage.json"

# nome_da_planilha -> rótulo exibido na interface
LABELS = {
    "socializacao": "Socialização",
    "linguagem": "Linguagem",
    "cognicao": "Cognição",
    "autocuidados": "Autocuidados",
    "des_mot": "Desenvolvimento Motor",
}


def main() -> None:
    wb = openpyxl.load_workbook(XLSX_FILE, data_only=True)
    data = {"areas": []}

    for nome in wb.sheetnames:
        ws = wb[nome]
        linhas = list(ws.iter_rows(values_only=True))
        itens = []
        for item, habilidade, range_i, range_f, _realiza in linhas[1:]:
            if habilidade is None or str(habilidade).strip() == "":
                continue
            itens.append(
                {
                    "item": int(item) if item is not None else None,
                    "habilidade": str(habilidade).strip(),
                    "range_i": int(range_i) if range_i is not None else 0,
                    "range_f": int(range_f) if range_f is not None else None,
                }
            )
        data["areas"].append(
            {"key": nome, "label": LABELS.get(nome, nome), "items": itens}
        )

    conteudo_json = json.dumps(data, ensure_ascii=False, indent=2)

    JSON_FILE.parent.mkdir(parents=True, exist_ok=True)
    JSON_FILE.write_text(conteudo_json + "\n", encoding="utf-8")

    total = sum(len(a["items"]) for a in data["areas"])
    print(
        f"Gerado {JSON_FILE} "
        f"com {len(data['areas'])} áreas e {total} habilidades."
    )


if __name__ == "__main__":
    main()
